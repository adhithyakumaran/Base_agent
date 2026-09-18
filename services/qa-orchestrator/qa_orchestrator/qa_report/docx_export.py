"""DOCX export from canonical QAReport."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

from qa_orchestrator.qa_report.models import QAReport


def render_docx_bytes(report: QAReport, *, repo_root: Path | None = None) -> bytes:
    try:
        from docx import Document
        from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
        from docx.shared import Inches
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("python-docx is not installed. pip install python-docx") from exc

    doc = Document()
    title = doc.add_heading("ScoutAI Enterprise QA Execution Report", 0)
    title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    doc.add_paragraph(f"Run: {report.run_id}")
    doc.add_paragraph(f"Goal: {report.goal}")
    doc.add_paragraph(f"Result: {report.overall_result}")
    doc.add_paragraph(f"Generated: {report.generated_at}")

    doc.add_heading("Executive summary", level=1)
    doc.add_paragraph(report.executive_summary)

    doc.add_heading("Test details", level=1)
    table = doc.add_table(rows=1, cols=2)
    hdr = table.rows[0].cells
    hdr[0].text = "Field"
    hdr[1].text = "Value"
    rows = [
        ("Flow IDs", ", ".join(report.flow_ids)),
        ("Parameters", str(report.input_parameters)),
        ("Environment", report.environment),
        ("Execution mode", str(report.execution_mode)),
        ("Command", str(report.execution.command)),
    ]
    for key, val in rows:
        cells = table.add_row().cells
        cells[0].text = key
        cells[1].text = val

    doc.add_heading("Validation", level=1)
    doc.add_paragraph(f"Phase: {report.validation.phase}")
    doc.add_paragraph(f"Result: {report.validation.validation_result}")
    doc.add_paragraph(f"Reason code: {report.validation.reason_code}")
    doc.add_paragraph(f"Failed checks: {', '.join(report.validation.failed_checks)}")

    doc.add_heading("Final result", level=1)
    if report.final_result:
        doc.add_paragraph(report.final_result.explanation)

    doc.add_heading("Evidence", level=1)
    for ev in report.evidence.screenshots[:20]:
        doc.add_paragraph(f"{ev.evidence_id}: {ev.description or ev.action or 'capture'} — {ev.path}")
        if repo_root:
            src = (repo_root / ev.path).resolve()
            if src.is_file():
                try:
                    doc.add_picture(str(src), width=Inches(5.5))
                except Exception:
                    pass

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
